import json
from typing import ClassVar
import re

from numpy import mean
from docx.table import Table as TableDocx
from bs4 import BeautifulSoup

from .util import is_number, drop_space_around_punctuation, normalize_spaces, is_bold
from .Cell import Cell
from .Item import Item
from .TableType import TableType
from .Paragraph import INCLUDE_XML, INDENT


NOTE_PATTERN = re.compile(r'([*]+)\s+([^*]+[^*\s])')


def get_aligned_cell(cells: list[Cell], col_offset: int):
    n_cols = 0

    for cell in cells:
        if col_offset <= n_cols:
            return cell

        n_cols += cell.n_cols

    return cell


class Table(Item):
    type_label: ClassVar[str] = 'table'

    def __init__(self, soup: BeautifulSoup, rows: list[list[Cell]], label: str):
        self.rows = rows
        self.soup = soup
        self.label = label

        self.contexts = None
        self.title = None
        self.id = None
        self.kind = None

        self._stats = None

    @classmethod
    def from_normacs_json(cls, json: dict):
        rows = []

        offset_to_cell_id = {}
        cell_id_to_n_remaining_rows = {}
        cell_id_to_cell = {}

        for row in json['rows']:
            row_ = []
            offset = 0

            for cell in row['cells']:
                while True:
                    if (spanning_cell_id := offset_to_cell_id.get(offset)) is not None:
                        spanning_cell = cell_id_to_cell.get(spanning_cell_id)

                        row_.append(spanning_cell.make_placeholder())

                        n_remaining_rows = cell_id_to_n_remaining_rows.get(spanning_cell_id)

                        if n_remaining_rows > 1:
                            cell_id_to_n_remaining_rows[spanning_cell_id] = n_remaining_rows - 1
                        else:
                            cell_id_to_n_remaining_rows.pop(spanning_cell_id)
                            cell_id_to_cell.pop(spanning_cell_id)
                            offset_to_cell_id.pop(offset)

                        offset += spanning_cell.n_cols
                    else:
                        break

                inlines = []

                cell_format = cell['cellFormat']

                n_rows = cell_format['rowSpan']
                n_cols = cell_format['columnSpan']

                for block in cell['blocks']:
                    for inline in block['inlines']:
                        if 'name' in inline:  # picture
                            pass
                            # inlines.append(inline['name'])
                        else:
                            inlines.append(inline['text'])

                row_.append(
                    cell := Cell(
                        drop_space_around_punctuation(
                            normalize_spaces(
                                ' '.join(inlines)
                            )
                        ),
                        n_rows = n_rows,
                        n_cols = n_cols
                    )
                )

                if n_rows > 1:
                    cell_id_to_cell[cell.id] = cell
                    cell_id_to_n_remaining_rows[cell.id] = n_rows - 1
                    offset_to_cell_id[offset] = cell.id

                offset += n_cols

            rows.append(row_)

        # for row in rows:
        #     print(row)

        return cls(
            soup = None,
            rows = rows,
            label = None
        )

    @classmethod
    def from_json(cls, json: dict, make_context: callable, label: str = None):
        if label is None:
            label = json.get('label')

        xml = json.get('xml')
        rows = Cell.deserialize_rows(json['rows'])
        table = cls(BeautifulSoup(xml, 'lxml'), rows = rows, label = label)

        table.id = json.get('id')
        table.kind = TableType(json.get('kind'))  # TODO: Rename TableType to TableKind
        table.title = json.get('title')
        table.contexts = [
            make_context(context)
            for context in json.get('contexts')
        ]

        return table

    def set_title(self, text: str = None):
        bold_text = [] if text is None else None

        if bold_text is not None:
            for fragment in self.soup.find_all('w:r'):
                if is_bold(fragment):
                    bold_text.append(fragment.text)

            text = normalize_spaces(' '.join(bold_text))

            if text.endswith(':'):
                text = text[:-1]

        self.title = text

    @classmethod
    def from_soup(cls, soup: BeautifulSoup, label: str):
        rows = []
        last_row = None

        for row in soup.find_all('w:tr'):
            cells = []

            col_offset = 0

            for cell in row.find_all('w:tc'):

                if last_row is not None and (vertical_span := cell.find('w:vmerge')) is not None and vertical_span.get('w:val') != 'restart':
                    cells.append(
                        placeholder := get_aligned_cell(last_row, col_offset).make_placeholder()
                    )

                    (origin := placeholder.origin).n_rows += 1
                    col_offset += origin.n_cols
                else:
                    n_cols = 1

                    if (horizontal_span := cell.find('w:gridspan')) is not None and (span_size := horizontal_span.get('w:val')) is not None:
                        n_cols = int(span_size)

                    col_offset += n_cols

                    cells.append(
                        Cell(
                            drop_space_around_punctuation(
                                normalize_spaces(cell.text)
                            ),
                            n_cols = n_cols
                        )
                    )

            last_row = cells
            rows.append(cells)

        # Parse notes - just add the note text in brackets after the original cell content without removing the anchor symbol(s)

        notes = {}

        # if rows[0][0].text.startswith('Сооружения'):

        for row in rows:
            for cell in row:
                if cell.text is not None:
                    for note in NOTE_PATTERN.findall(cell.text):
                        notes[note[0]] = note[1]

        for row in rows:
            for cell in row:
                for key in sorted(notes.keys(), key = lambda key: len(key), reverse = True):
                    if cell.text is not None and cell.text.endswith(key):
                        # cell.text = f'{cell.text.strip()[:-(len(key) + 2)]} ({value})'
                        cell.text = normalize_spaces(f'{cell.text} ({notes[key]})')

        return cls(soup, rows, label)

    def to_json(self, path: str = None, indent: int = INDENT):
        data = Cell.serialize_rows(self.rows)

        data['label'] = self.label

        if INCLUDE_XML:
            data['xml'] = str(self.soup)

        if (contexts := self.contexts) is not None:
            data['contexts'] = [context.json for context in contexts]

        if (title := self.title) is not None:
            data['title'] = title

        if (id_ := self.id) is not None:
            data['id'] = id_

        if (kind := self.kind) is not None:
            data['kind'] = kind.value

        data['type'] = self.type_label

        if path is not None:
            with open(path, 'w', encoding = 'utf-8') as file:
                json.dump(data, file, indent = indent, ensure_ascii = False)

        return data

    @property
    def json(self):
        return self.to_json()

    @property
    def n_chars(self):
        return len(json.dumps(self.json, ensure_ascii = False))

    @property
    def stats(self):
        if self._stats is None:
            self._stats = TableStats(self)

        return self._stats

    @property
    def content(self):
        return [
            [
                cell.text
                for cell in row
            ]
            for row in self.rows
        ]

    @property
    def next_sibling_paragraphs(self):
        return self.soup.findNextSiblings('w:p')

    @property
    def previous_sibling_paragraphs(self):
        return self.soup.findPreviousSiblings('w:p')

    @property
    def next_sibling_paragraph(self):
        for item in self.next_sibling_paragraphs:
            if (text := item.text) is not None and len(text.strip()) > 0:
                return item

        return None

    @property
    def previous_sibling_paragraph(self):
        for item in self.previous_sibling_paragraphs:
            if (text := item.text) is not None and len(text.strip()) > 0:
                return item

        return None

    @property
    def isotropic(self):
        row_length = None

        for row in self.rows:
            if row_length is None:
                row_length = len(row)
                continue

            if len(row) != row_length or any(cell.text is None for cell in row):
                return False

        return True

    @property
    def as_text(self):
        if not self.isotropic:
            raise ValueError("Can't convert non-isotropic table into text")

        lines = []  # if (title := self.title) is None else [title]

        for row in self.rows:
            lines.append(' '.join(cell.text for cell in row))

        return '\n'.join(lines)

    @property
    def as_texts(self):
        return [
            cell.text
            for row in self.rows
            for cell in row
        ]


class TableStats:
    def __init__(self, table: Table):
        self._item = table

        n_cells = 0
        n_rows = 0
        n_cols = 0
        n_chars = []

        n_numeric_cells = 0
        n_text_cells = 0
        n_empty_cells = 0

        n_chars_numeric = []
        n_chars_text = []

        n_rowspans = []
        n_colspans = []

        for row in table.data['rows']:
            n_rows += 1

            n_cols_current_row = 0
            n_cols_current_row_without_placeholders = 0

            for cell in row:
                n_cols_current_row += 1

                if (cell_text := cell.get('text')) is not None:
                    n_cols_current_row_without_placeholders += 1
                    n_chars.append(cell_length := len(cell_text))

                    if is_number(cell_text):
                        n_numeric_cells += 1
                        n_chars_numeric.append(cell_length)
                    elif len(cell_text) < 1:
                        n_empty_cells += 1
                    else:
                        n_text_cells += 1
                        n_chars_text.append(cell_length)

                    n_rowspans.append(cell.get('rows'))
                    n_colspans.append(cell.get('cols'))

            n_cells += n_cols_current_row_without_placeholders

            if n_cols_current_row > n_cols:
                n_cols = n_cols_current_row

        # print(n_cells)

        self.n_cells = n_cells
        self.n_rows = n_rows
        self.n_cols = n_cols
        self.n_chars = n_chars

        self.n_numeric_cells = n_numeric_cells
        self.n_chars_numeric = n_chars_numeric

        self.n_text_cells = n_text_cells
        self.n_chars_text = n_chars_text
        self.n_empty_cells = n_empty_cells

        self.n_rowspans = n_rowspans
        self.n_colspans = n_colspans

    @property
    def as_vector(self):
        return [
            self.n_cells,
            self.n_rows,
            self.n_cols,
            sum(self.n_chars), 0 if len(self.n_chars) < 1 else mean(self.n_chars),
            self.n_numeric_cells,
            self.n_text_cells,
            self.n_empty_cells,
            sum(self.n_chars_numeric), 0 if len(self.n_chars_numeric) < 1 else mean(self.n_chars_numeric),
            sum(self.n_chars_text), 0 if len(self.n_chars_text) < 1 else mean(self.n_chars_text),
            self.mean_rowspan,
            self.mean_colspan
        ]

    @classmethod
    @property
    def vector_legend(self):
        return [
            'n-cells',
            'n-rows',
            'n-cols',
            'n-chars-sum', 'n-chars-mean',
            'n-numeric-cells',
            'n-text-cells',
            'n-empty-cells',
            'n-chars-numeric-sum', 'n-chars-numeric-mean',
            'n-chars-text-sum', 'n-chars-text-mean',
            'mean-rowspan',
            'mean-colspan'
        ]

    @property
    def mean_rowspan(self):
        return 0 if len(self.n_rowspans) < 1 else mean(self.n_rowspans)

    @property
    def mean_colspan(self):
        return 0 if len(self.n_colspans) < 1 else mean(self.n_colspans)
